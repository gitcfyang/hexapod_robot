#!/usr/bin/env python3
"""
xlsx BOM → csv + md 转换器 (提取超链接)

问题: xlsx 单元格可携带超链接 (显示文本 + URL 目标), 但 CSV 格式
无超链接概念, Excel/WPS 另存 CSV 时只写显示文本, URL 丢失。

本脚本用 openpyxl 读取每个单元格的链接目标 (cell.hyperlink.target):
  - 生成 hexapod_bom.csv  — "链接"列填 URL 全文 (纯文本, 任何工具可读)
  - 生成 hexapod_bom.md   — markdown 表格, GitHub 网页上链接可直接点击

用法:
  pip install openpyxl
  python3 tools/xlsx_bom_to_csv.py <bom.xlsx> <输出目录>

默认输入 hardware/bom/hexapod_bom.xlsx, 输出到同目录。
xlsx 重新编辑后重新运行本脚本刷新 csv/md 即可 (xlsx 不入库)。
"""

import sys
import csv
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    sys.exit("需要 openpyxl: pip install openpyxl")


def cell_text(cell):
    """单元格文本: 优先链接目标 URL, 无链接则取显示值"""
    if cell.hyperlink and cell.hyperlink.target:
        return cell.hyperlink.target.strip()
    v = cell.value
    if v is None:
        return ""
    return str(v).strip()


def main():
    src = Path(sys.argv[1]) if len(sys.argv) > 1 else \
        Path("hardware/bom/hexapod_bom.xlsx")
    out_dir = Path(sys.argv[2]) if len(sys.argv) > 2 else src.parent
    if not src.exists():
        sys.exit(f"找不到 {src} — 请把原 xlsx 放回该路径")

    wb = load_workbook(src, data_only=True)
    ws = wb.active

    rows = [[cell_text(c) for c in row] for row in ws.iter_rows()]
    # 去掉全空行
    rows = [r for r in rows if any(r)]
    if not rows:
        sys.exit("表格为空")
    ncol = max(len(r) for r in rows)
    rows = [r[:ncol] for r in rows]
    # 去掉尾部全空列 (xlsx 常残留空白列)
    while ncol > 1 and all(len(r) <= ncol - 1 or not r[ncol - 1] for r in rows):
        ncol -= 1
        rows = [r[:ncol] for r in rows]

    # CSV (utf-8-sig 带 BOM, Excel/WPS 直接打开中文不乱码)
    csv_path = out_dir / "hexapod_bom.csv"
    with open(csv_path, "w", newline="", encoding="utf-8-sig") as f:
        csv.writer(f).writerows(rows)
    print(f"已生成 {csv_path}")

    # Markdown 表格 (链接可点击)
    header, data = rows[0], rows[1:]
    md_path = out_dir / "hexapod_bom.md"
    lines = ["# 全项目物料清单 (BOM)", ""]
    lines.append("| " + " | ".join(header) + " |")
    lines.append("|" + "---|" * len(header))
    for r in data:
        cells = []
        for c in r:
            if c.startswith("http"):
                cells.append(f"[🔗]({c})")
            else:
                cells.append(c.replace("|", "\\|") or " ")
        lines.append("| " + " | ".join(cells) + " |")
    lines.append("")
    lines.append("> 由 `tools/xlsx_bom_to_csv.py` 从 xlsx 生成 (xlsx 为本地编辑源, 不入库)。")
    lines.append("> 修改物料: 直接编辑 csv 的链接列 (URL 文本), 或保留本地 xlsx 重新生成。")
    (out_dir / md_path.name).write_text("\n".join(lines), encoding="utf-8")
    print(f"已生成 {md_path}")


if __name__ == "__main__":
    main()
